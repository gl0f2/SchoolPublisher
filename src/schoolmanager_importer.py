"""Import der aktuellen Schoolmanager-Unterrichtsdaten (A/B-Wochen)."""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
import re

from openpyxl import load_workbook

from models import Unterricht


class SchoolmanagerImporter:
    """
    Liest die von schoolmanager_import_v03_ab.py erzeugten A/B-Wochen ein.

    Pro Kombination (Klasse, Fach, Lehrkraft) wird GENAU EIN Unterricht-Objekt
    erzeugt. Unterschiedliche studentGroups werden als Kopplungsinformation
    zusammengeführt und erzeugen keine doppelten Unterrichtseinträge.
    """

    IGNORIERTE_KLASSEN = {"AG", "Frei", "Fö", "KOOP_12", "ORG_13", "xy"}

    def __init__(self, datenverzeichnis: Path):
        self.datenverzeichnis = Path(datenverzeichnis)
        self.stundenplanverzeichnis = self.datenverzeichnis / "stundenplaene_ab"

    def load(self) -> list[Unterricht]:
        if not self.stundenplanverzeichnis.is_dir():
            raise FileNotFoundError(
                f"Schoolmanager-Stundenplandaten nicht gefunden: "
                f"{self.stundenplanverzeichnis}"
            )

        unterrichtsliste = []

        for klassenordner in sorted(self.stundenplanverzeichnis.iterdir()):
            if not klassenordner.is_dir():
                continue

            klasse = klassenordner.name
            if klasse in self.IGNORIERTE_KLASSEN:
                continue

            datei_a = klassenordner / "woche_A.json"
            datei_b = klassenordner / "woche_B.json"

            if not datei_a.is_file() or not datei_b.is_file():
                continue

            woche_a = self._woche_auswerten(klasse, self._json_laden(datei_a))
            woche_b = self._woche_auswerten(klasse, self._json_laden(datei_b))

            alle_schluessel = set(woche_a) | set(woche_b)

            for key in alle_schluessel:
                a = woche_a.get(key)
                b = woche_b.get(key)
                info = a or b

                stunden_a = len(a["termine"]) if a else 0
                stunden_b = len(b["termine"]) if b else 0
                wochenstunden = (stunden_a + stunden_b) / 2.0

                gruppen = set()
                if a:
                    gruppen.update(a["gruppen"])
                if b:
                    gruppen.update(b["gruppen"])

                kopplung = ", ".join(sorted(gruppen, key=str.casefold)) or None

                unterrichtsliste.append(
                    Unterricht(
                        klasse=info["klasse"],
                        fach=info["fach"],
                        fachname=info["fachname"],
                        lehrer=info["lehrer"],
                        wochenstunden=wochenstunden,
                        stundenplan_name=info["stundenplan_name"],
                        kopplung=kopplung,
                    )
                )

        unterrichtsliste = self._manuelle_lehrerwechsel_anwenden(unterrichtsliste)
        unterrichtsliste = self._halbjahresdaten_anwenden(unterrichtsliste)

        return sorted(
            unterrichtsliste,
            key=lambda e: (
                e.klasse.casefold(),
                e.anzeigename.casefold(),
                e.lehrer.casefold(),
            ),
        )


    def _halbjahresdaten_anwenden(
        self, unterrichtsliste: list[Unterricht]
    ) -> list[Unterricht]:
        """Ergänzt 1./2.-Halbjahresunterricht aus ``Daten JND.xlsx``.

        Schoolmanager bleibt die Hauptquelle. Die Exceldatei liefert nur die
        Halbjahreszuordnung und ggf. eine Lehrkraft des zweiten Halbjahres,
        die im aktuellen September-Stundenplan noch nicht vorkommt.
        """
        projektwurzel = self.datenverzeichnis.parent.parent
        pfad = projektwurzel / "Daten JND.xlsx"
        if not pfad.is_file():
            return unterrichtsliste

        wb = load_workbook(pfad, data_only=True)
        ws = wb.active
        halbjahreszeilen: list[dict] = []

        try:
            for zeile in range(5, ws.max_row + 1):
                lehrer = str(ws.cell(zeile, 5).value or "").strip()
                fach = str(ws.cell(zeile, 6).value or "").strip()
                klassenwert = str(ws.cell(zeile, 7).value or "").strip()
                wst_wert = ws.cell(zeile, 4).value

                marker = {
                    str(ws.cell(zeile, 13).value or "").strip().casefold(),
                    str(ws.cell(zeile, 14).value or "").strip().casefold(),
                }
                if "1.hj" in marker:
                    hj = 1
                elif "2.hj" in marker:
                    hj = 2
                else:
                    continue

                if not fach or not lehrer:
                    continue

                fach_key = fach.casefold().replace("–", "-").replace("—", "-").replace("‑", "-")
                if (
                    fach_key in {"kl", "kl2", "ch-pr", "b-pr", "ph-pr"}
                    or fach_key.startswith("mug-")
                    or fach_key.startswith("spj")
                ):
                    continue

                try:
                    wst = float(str(wst_wert).replace(",", ".")) if wst_wert not in (None, "") else 0.0
                except (TypeError, ValueError):
                    wst = 0.0

                for teil in klassenwert.split(","):
                    klasse = teil.strip().lower()
                    if not re.fullmatch(r"(?:5|6|7|8|9|10)[a-f]", klasse):
                        continue

                    fach_norm, fachname = self._fach_normalisieren(klasse, fach, self._fachname_fuer_excelcode(fach))
                    if fach_norm is None:
                        continue

                    halbjahreszeilen.append({
                        "klasse": klasse,
                        "fach": fach_norm,
                        "fachname": fachname,
                        "lehrer": lehrer,
                        "halbjahr": hj,
                        "wochenstunden": wst,
                    })
        finally:
            wb.close()

        # Doppelte Excelzeilen entfernen.
        eindeutig: dict[tuple[str, str, str, int], dict] = {}
        for e in halbjahreszeilen:
            key = (e["klasse"].casefold(), e["fach"].casefold(), e["lehrer"].casefold(), e["halbjahr"])
            eindeutig[key] = e
        halbjahreszeilen = list(eindeutig.values())

        # Derselbe Lehrer in beiden Halbjahren = ganzjährig.
        hj_pro_lehrer: dict[tuple[str, str, str], set[int]] = defaultdict(set)
        for e in halbjahreszeilen:
            hj_pro_lehrer[(e["klasse"].casefold(), e["fach"].casefold(), e["lehrer"].casefold())].add(e["halbjahr"])

        excel_faecher = {(e["klasse"].casefold(), e["fach"].casefold()) for e in halbjahreszeilen}

        # Für Fächer mit Halbjahresangaben ersetzt die Excel-Zuordnung die
        # aktuelle Lehrerzuordnung dieses Faches. Alle anderen Fächer bleiben
        # unverändert aus Schoolmanager erhalten.
        basis_pro_fach: dict[tuple[str, str], Unterricht] = {}
        ergebnis: list[Unterricht] = []
        for u in unterrichtsliste:
            key = (u.klasse.casefold(), u.fach.casefold())
            basis_pro_fach.setdefault(key, u)
            if key not in excel_faecher:
                ergebnis.append(u)

        for e in halbjahreszeilen:
            lehrer_key = (e["klasse"].casefold(), e["fach"].casefold(), e["lehrer"].casefold())
            # Wenn beide Halbjahre denselben Lehrer nennen, nur einmal ganzjährig erzeugen.
            if hj_pro_lehrer[lehrer_key] == {1, 2} and e["halbjahr"] == 2:
                continue
            hj = None if hj_pro_lehrer[lehrer_key] == {1, 2} else e["halbjahr"]

            basis = basis_pro_fach.get((e["klasse"].casefold(), e["fach"].casefold()))
            ergebnis.append(Unterricht(
                klasse=basis.klasse if basis else e["klasse"],
                fach=basis.fach if basis else e["fach"],
                fachname=basis.fachname if basis else e["fachname"],
                lehrer=e["lehrer"],
                wochenstunden=(basis.wochenstunden if basis and basis.wochenstunden > 0 else e["wochenstunden"]),
                stundenplan_name=basis.stundenplan_name if basis else e["fachname"],
                kopplung=basis.kopplung if basis else None,
                halbjahr=hj,
            ))

        return ergebnis

    @staticmethod
    def _fachname_fuer_excelcode(fach: str) -> str:
        """Lesbarer Fachname für Halbjahresfächer, die im aktuellen Plan fehlen."""
        key = fach.strip().casefold()
        namen = {
            "d": "Deutsch", "m": "Mathematik", "e": "Englisch",
            "b": "Biologie", "bk": "Bildende Kunst", "mu": "Musik",
            "gk": "Gemeinschaftskunde", "geo": "Geographie", "wbs": "WBS",
            "ium": "Informatik und Medienbildung", "nit": "NIT",
            "stern": "Sternstunde",
        }
        return namen.get(key, fach.strip())

    def _manuelle_lehrerwechsel_anwenden(
        self, unterrichtsliste: list[Unterricht]
    ) -> list[Unterricht]:
        """Wendet lokale Lehrerwechsel nach dem Schoolmanager-Import an.

        Die Datei ``manuelle_aenderungen.json`` bleibt unabhängig von den
        automatisch aktualisierten A/B-Stundenplandaten erhalten.
        """
        pfad = self.datenverzeichnis / "manuelle_aenderungen.json"
        if not pfad.is_file():
            return unterrichtsliste

        with pfad.open("r", encoding="utf-8") as f:
            daten = json.load(f)

        wechsel = daten.get("lehrerwechsel", []) if isinstance(daten, dict) else []
        if not isinstance(wechsel, list):
            raise ValueError(f"'lehrerwechsel' muss eine Liste sein: {pfad}")

        ergebnis = list(unterrichtsliste)
        for eintrag in wechsel:
            if not isinstance(eintrag, dict):
                continue
            klasse = str(eintrag.get("klasse") or "").strip()
            fach = str(eintrag.get("fach") or "").strip().casefold()
            lehrer = str(eintrag.get("lehrer") or "").strip()
            if not klasse or not fach or not lehrer:
                raise ValueError(f"Unvollständiger Lehrerwechsel in {pfad}: {eintrag}")

            treffer = []
            for i, u in enumerate(ergebnis):
                fachwerte = {u.fach.strip().casefold(), u.fachname.strip().casefold()}
                if u.klasse.casefold() == klasse.casefold() and fach in fachwerte:
                    treffer.append(i)

            if not treffer:
                raise ValueError(
                    f"Manueller Lehrerwechsel nicht gefunden: {klasse} / {eintrag.get('fach')}"
                )

            # Ein Lehrerwechsel ersetzt alle bisher aus Schoolmanager gelesenen
            # Lehrkräfte für genau diese Klasse/Fach-Kombination. Bei mehreren
            # alten Einträgen wird nur ein Unterrichtsobjekt übernommen.
            basis = ergebnis[treffer[0]]
            neu = Unterricht(
                klasse=basis.klasse,
                fach=basis.fach,
                fachname=basis.fachname,
                lehrer=lehrer,
                wochenstunden=basis.wochenstunden,
                stundenplan_name=basis.stundenplan_name,
                kopplung=basis.kopplung,
            )
            ergebnis = [u for i, u in enumerate(ergebnis) if i not in set(treffer)]
            ergebnis.append(neu)

        return ergebnis

    def _woche_auswerten(self, klasse: str, lessons: list) -> dict:
        """
        Gruppiert ausschließlich nach (Klasse, Fach, Lehrer).

        Die Schülergruppen werden nur gesammelt. Für die Stundenzahl werden
        eindeutige (Datum, Unterrichtsstunde)-Termine gezählt. Dadurch erzeugt
        dieselbe Stunde mit mehreren studentGroups keine Mehrfachzählung.
        """
        result = {}

        for lesson in lessons:
            if not isinstance(lesson, dict):
                continue

            # SchoolPublisher bildet die reguläre Unterrichtsverteilung ab.
            # Vertretungen/Änderungen (changedLesson) dürfen deshalb keine
            # zusätzlichen Fachlehrkräfte auf den Karteikarten erzeugen.
            if lesson.get("type") != "regularLesson":
                continue

            actual = lesson.get("actualLesson")
            if not isinstance(actual, dict):
                continue

            # In Klassen 10 sind die 7. und 8. Stunde für das Blatt nicht relevant.
            if klasse.startswith("10"):
                nummer = (lesson.get("classHour") or {}).get("number")
                try:
                    if int(nummer) in {7, 8}:
                        continue
                except (TypeError, ValueError):
                    pass

            subject = actual.get("subject") or {}
            fach = self._subject_code(subject)
            fachname = self._subject_name(subject)
            if not fach:
                continue

            # MPG-Sonderregeln für die Fachbezeichnungen im Stundenplan.
            fach, fachname = self._fach_normalisieren(klasse, fach, fachname)
            if fach is None:
                # z. B. Biologie-/Physik-Praktika: keine eigene Karteikarte.
                continue

            gruppen = self._student_groups(actual)
            termin = self._termin_key(lesson)

            for teacher in actual.get("teachers") or []:
                if not isinstance(teacher, dict):
                    continue

                kuerzel = str(teacher.get("abbreviation") or "").strip()
                if not kuerzel:
                    continue

                key = (
                    klasse.casefold(),
                    fach.casefold(),
                    kuerzel.casefold(),
                )

                if key not in result:
                    result[key] = {
                        "klasse": klasse,
                        "fach": fach,
                        "fachname": fachname,
                        "lehrer": kuerzel,
                        "stundenplan_name": fachname or fach,
                        "gruppen": set(),
                        "termine": set(),
                    }

                result[key]["gruppen"].update(gruppen)
                result[key]["termine"].add(termin)

        return result


    @staticmethod
    def _fach_normalisieren(klasse: str, fach: str, fachname: str) -> tuple[str | None, str]:
        code = fach.strip()
        name = fachname.strip()
        code_key = code.casefold().replace("–", "-").replace("—", "-").replace("‑", "-")
        text_key = f"{code} {name}".casefold()
        stufe_text = "".join(z for z in klasse if z.isdigit())
        stufe = int(stufe_text) if stufe_text else 0

        # Klassen 8-10: SPO = Profilfach, SPO-M/SPO-W = regulärer Sport.
        if 8 <= stufe <= 10:
            if code_key == "spo":
                return "SPO-PROFIL", "Sport (Profil)"
            if code_key in {"spo-m", "spo-w"}:
                return "SPO", "Sport"

        # Klassenlehrerstunden und Gesangsklassen erzeugen keine eigene Karte.
        # Sternstunde und MTW werden dagegen für die Elternabendübersicht benötigt.
        if (
            code_key in {"klassenlehrer", "kl"}
            or code_key.startswith("mug-")
        ):
            return None, name

        # Sternstunde:
        # Klasse 6 -> "Sternstunde"
        # Klasse 7 -> "Sternstunde/Mentoring"
        if "sternstunde" in text_key or code_key in {"stern", "sternstunde"}:
            if stufe == 7:
                return "Stern", "Sternstunde/Mentoring"
            return "Stern", "Sternstunde"

        # MTW wird in Klasse 5 als eigene Karte angezeigt.
        if code_key == "mtw":
            return "MTW", "Musik- und Theaterwerkstatt"

        # Praktika in Biologie, Physik und Chemie werden nicht dargestellt.
        if "prakt" in text_key and any(w in text_key for w in (
            "bio", "biologie", "phys", "physik", "chem", "chemie"
        )):
            return None, name

        return code, name

    @staticmethod
    def _termin_key(lesson: dict):
        class_hour = lesson.get("classHour") or {}
        nummer = class_hour.get("number")
        # Fallback auf IDs/Zeiten, falls number einmal fehlen sollte.
        if nummer is None:
            nummer = (
                class_hour.get("id")
                or class_hour.get("startTime")
                or class_hour.get("start")
                or json.dumps(class_hour, sort_keys=True, ensure_ascii=False)
            )
        return (lesson.get("date"), str(nummer))

    @staticmethod
    def _json_laden(path: Path) -> list:
        with path.open("r", encoding="utf-8") as f:
            daten = json.load(f)
        if not isinstance(daten, list):
            raise ValueError(f"Erwartete Liste in {path}")
        return daten

    @staticmethod
    def _subject_code(subject: dict) -> str:
        return str(
            subject.get("abbreviation")
            or subject.get("shortName")
            or subject.get("name")
            or subject.get("id")
            or ""
        ).strip()

    @staticmethod
    def _subject_name(subject: dict) -> str:
        return str(
            subject.get("name")
            or subject.get("abbreviation")
            or subject.get("shortName")
            or subject.get("id")
            or ""
        ).strip()

    @staticmethod
    def _student_groups(actual: dict) -> set[str]:
        namen = set()
        for group in actual.get("studentGroups") or []:
            if not isinstance(group, dict):
                continue
            name = str(
                group.get("name")
                or group.get("abbreviation")
                or group.get("shortName")
                or ""
            ).strip()
            if name:
                namen.add(name)
        return namen
